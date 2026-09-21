#!/usr/bin/env python3
"""
Builds the ScoutLite player research brief as a .docx file.

Structure, v3 (2026-09-17 -- both Signals are now fully deterministic, no LLM involved in
either number/label, only in the prose explaining them):
  - Signals block at the top -- Quality (a label, e.g. "World Class", from percentile-ranked
    per-90 stats, possession-adjusted where relevant) and Fit (a label, e.g. "Hand-in-Glove
    Fit", from a percentile-profile comparison against a real reference club's current squad).
    Either can come back unavailable (uncovered league/position, or no philosophy given) --
    shown as "not available" rather than a fabricated result.
  - 1. Who he is -- bio/background (pure data, no LLM)
  - 2. Stats & performance -- season/misc/keeper/xG tables (pure data, no LLM)
  - 3. What people say -- news headlines (data, each linked to its actual article) + a short
    LLM-synthesized read
  - 4. Signals & fit read -- the Fit label + reference-club comparison, with the LLM's prose
    explaining it (not deciding it) + scout's role notes + philosophy chosen
  - 5. Sources -- every link a scout needs to check a claim themselves: the FBref profile the
    stats came from, the Understat profile the xG/xA came from, and NewsAPI's publisher/date
    per headline (also inline above) -- this is a research brief, not a black box
  - 6. Understanding the Signals -- a standing, identical-every-run glossary explaining what
    Quality and Fit actually measure and don't, so a scout unfamiliar with the tool isn't left
    guessing what a label means

Deliberately named "research brief," never "report" or "verdict" -- ScoutLite is a data/
research layer for a scout, not a conclusion reached on their behalf.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# Shown whenever a club philosophy was actually assessed. Added after the Track C eval
# (2026-09-16): 7 players, all 6 possible philosophy combinations, 21 runs -- every single one
# landed on Fit: 3/5 back when Fit was an LLM-judged number. ScoutLite has no pace, sprint, or
# pressing-volume data in any source, and that's exactly what philosophy fit depends on -- the
# model was doing the right thing by refusing to guess rather than fabricate confidence, but the
# result never actually moved. Fit is now a deterministic profile-similarity comparison instead
# (v3, 2026-09-17), which DOES vary meaningfully -- but the same underlying data gap is still
# real: the comparison itself never sees pace or pressing, only counting/creative stats, so even
# a close statistical match doesn't confirm tactical fit. See eval/TRACK_C_REPORT.md for the
# original evidence.
FIT_SCOPE_CAVEAT = (
    "Fit cannot assess pace, sprint, or pressing intensity -- no ScoutLite data source "
    "measures these, and they're central to what a club philosophy actually demands. Even a "
    "close statistical match above doesn't confirm tactical fit on those dimensions."
)

# For the "6. Understanding the Signals" glossary -- same cutoffs as scoring.QUALITY_LABELS,
# just phrased for a reader rather than a comparison operator.
QUALITY_LABEL_SCALE = [
    ("Below 20th percentile", "Below Rotation"),
    ("20th–40th percentile", "Depth Option"),
    ("40th–60th percentile", "Solid Starter"),
    ("60th–80th percentile", "Strong Starter"),
    ("80th percentile and above", "World Class"),
]


def _add_hyperlink(paragraph, url: str, text: str):
    """python-docx has no built-in hyperlink support -- this is the standard low-level recipe:
    register the URL as an external relationship on the paragraph's part, then hand-build the
    <w:hyperlink> run pointing at it, styled like a normal link (blue, underlined)."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    run.append(run_props)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _add_dict_table(doc: Document, data: dict, skip: tuple[str, ...] = ()):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in data.items():
        if value in (None, "") or key in skip:
            continue
        row = table.add_row().cells
        row[0].text = key.replace("_", " ").title()
        row[1].text = str(value)


def build_docx(data: dict, output_path: Path) -> Path:
    """data: the same dict shape scoutlite_combined.research_player() returns -- player_name,
    player_url, bio, stats, xg, articles, misc, keeper, scout_notes, philosophy, news_synthesis,
    fit_read, quality, fit_signal, judge. Takes one cohesive object instead of ~16 individual
    same-typed params (2026-09-21) -- the old signature had misc/keeper/bio as adjacent
    same-typed `dict | None` positional args with nothing to stop them being transposed by
    mistake at a call site; a single dict keyed by name removes that risk entirely."""
    player_name = data["player_name"]
    player_url = data.get("player_url")
    bio = data["bio"]
    stats = data["stats"]
    xg = data.get("xg")
    articles = data.get("articles") or []
    misc = data.get("misc")
    keeper = data.get("keeper")
    news_synthesis = data.get("news_synthesis")
    fit_read = data.get("fit_read")
    scout_notes = data.get("scout_notes")
    philosophy = data.get("philosophy")
    quality = data.get("quality")
    fit_signal = data.get("fit_signal")
    judge = data.get("judge")

    doc = Document()

    title = doc.add_heading(f"{player_name} — Player Research Brief", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = doc.add_paragraph()
    subtitle.add_run(
        f"Season: {stats.get('season', 'unknown')}  ·  A data/research layer for a scout to "
        "weigh -- not a scouting verdict."
    ).italic = True

    # --- Confidence warning (only if the automated judge didn't clear the threshold) ---------
    if judge and judge.get("confidence_warning"):
        warn = doc.add_paragraph()
        warn.add_run(
            f"⚠ CONFIDENCE WARNING — this brief's two written paragraphs scored "
            f"{judge['source_accuracy']}% on automated claim-grounding after "
            f"{judge['iterations']} revision pass(es), below the "
            f"{judge.get('threshold', 80)}% threshold. Check the flagged points below before "
            f"relying on the written sections; the data tables and signals are unaffected."
        ).bold = True
        for finding in judge.get("findings", []):
            doc.add_paragraph(finding, style="List Bullet")

    # --- Signals block (top) -----------------------------------------------------------
    doc.add_heading("Signals", level=1)
    quality_label = quality["label"] if quality else None
    fit_label = fit_signal["label"] if fit_signal else None
    quality_text = quality_label or "not available"
    if quality and quality["raw_label"] != quality["label"]:
        quality_text += f" (raw: {quality['raw_label']})"
    fit_text = f"{fit_label} vs. {fit_signal['reference_club']}" if fit_signal else "not available"
    p = doc.add_paragraph()
    p.add_run(f"Quality: {quality_text}  ·  Fit: {fit_text}").bold = True
    doc.add_paragraph().add_run(
        "Signals for the scout to weigh, never a conclusion the tool reaches on the scout's "
        "behalf. Both are fully deterministic -- non-AI, percentile-based -- not an LLM "
        "judgment; the LLM only writes the prose explaining each one, never decides the label. "
        "See \"6. Understanding the Signals\" below for what each actually measures."
    ).italic = True
    if quality:
        doc.add_paragraph().add_run(quality["explanation"]).italic = True
        doc.add_paragraph().add_run(
            f"Exact breakdown ({quality['avg_percentile']} percentile average): " + ", ".join(
                f"{k.replace('_', ' ')} = {v:.0f} percentile" for k, v in quality["components"].items()
            )
        ).italic = True
        if quality["raw_avg_percentile"] != quality["avg_percentile"]:
            doc.add_paragraph().add_run(
                f"Without the possession adjustment: {quality['raw_avg_percentile']} percentile "
                f"average ({quality['raw_label']}) -- see \"6. Understanding the Signals\" for why "
                "these two numbers can differ."
            ).italic = True
        if quality.get("specialist_caveat"):
            doc.add_paragraph().add_run(quality["specialist_caveat"]).italic = True
    else:
        doc.add_paragraph().add_run(
            "Quality signal not available -- either this player's league isn't one of the 5 "
            "covered (Premier League, La Liga, Bundesliga, Serie A, Ligue 1), or there wasn't "
            "enough data for their position group this season."
        ).italic = True

    # --- 1. Who he is --------------------------------------------------------------------
    doc.add_heading("1. Who He Is", level=1)
    if bio:
        _add_dict_table(doc, bio)
    else:
        doc.add_paragraph("No background data available.")

    # --- 2. Stats & performance ------------------------------------------------------------
    doc.add_heading("2. Stats & Performance", level=1)
    doc.add_heading(f"Season stats ({stats.get('season', 'unknown')})", level=2)
    _add_dict_table(doc, stats)

    if keeper:
        doc.add_heading("Goalkeeping stats", level=2)
        _add_dict_table(doc, keeper)

    if misc:
        doc.add_heading("Defensive/discipline stats", level=2)
        _add_dict_table(doc, misc)

    if xg:
        doc.add_heading("Advanced stats (Understat: xG/xA)", level=2)
        doc.add_paragraph().add_run(
            f"⚠ Matched by name lookup (not a guaranteed-unique ID) to Understat's "
            f"\"{xg['understat_matched_name']}\" -- for compound surnames, nicknames, or "
            f"abbreviations, this can occasionally match the wrong player or miss a real one. "
            f"Verify the matched name against who you actually mean before trusting these numbers."
        ).italic = True
        if xg.get("understat_url"):
            link_p = doc.add_paragraph()
            link_p.add_run("Source: ")
            _add_hyperlink(link_p, xg["understat_url"], xg["understat_url"])
        _add_dict_table(doc, xg, skip=("understat_url",))
    else:
        doc.add_paragraph("Advanced stats (xG/xA): not available -- Understat doesn't cover this player's league, or no name match was found.")

    # --- 3. What people say --------------------------------------------------------------
    doc.add_heading("3. What People Say", level=1)
    doc.add_paragraph(news_synthesis or "No recent news synthesis available.")
    if articles:
        doc.add_heading("Recent headlines (last 28 days)", level=2)
        for a in articles[:8]:
            bullet = doc.add_paragraph(style="List Bullet")
            title = a.get("title", "")
            if a.get("url"):
                _add_hyperlink(bullet, a["url"], title)
            else:
                bullet.add_run(title)
            source = a.get("source", {}).get("name", "")
            published = (a.get("publishedAt") or "")[:10]
            attribution = " · ".join(v for v in (source, published) if v)
            if attribution:
                bullet.add_run(f"  ({attribution})").italic = True

    # --- 4. Signals & fit read -------------------------------------------------------------
    doc.add_heading("4. Signals & Fit Read", level=1)
    has_philosophy = philosophy and (philosophy.get("in_possession") or philosophy.get("out_of_possession"))
    if has_philosophy:
        style_desc = " / ".join(v for v in philosophy.values() if v)
        if fit_signal:
            doc.add_paragraph().add_run(
                f"Club philosophy assessed against: {style_desc}  ·  "
                f"Fit: {fit_signal['label']} vs. {fit_signal['reference_club']}"
            ).bold = True
            doc.add_paragraph().add_run(
                f"Exact comparison ({fit_signal['reference_club']}'s current squad, same "
                "position group): " + ", ".join(
                    f"{k.replace('_', ' ')} = this player {v:.0f} vs. "
                    f"{fit_signal['reference_components'][k]:.0f}"
                    for k, v in fit_signal["target_components"].items()
                )
            ).italic = True
        else:
            doc.add_paragraph().add_run(
                f"Club philosophy assessed against: {style_desc}  ·  Fit: not available"
            ).bold = True
            doc.add_paragraph().add_run(
                "Fit couldn't be computed for this player -- their league or position isn't "
                "covered by the reference-club comparison."
            ).italic = True
        doc.add_paragraph().add_run(FIT_SCOPE_CAVEAT).italic = True
    if scout_notes and scout_notes.strip():
        doc.add_paragraph().add_run(f"Scout's role notes: \"{scout_notes.strip()}\"").italic = True
    doc.add_paragraph(fit_read or "No fit signal available.")

    # --- 5. Sources -------------------------------------------------------------------------
    # Every link a scout would need to check a claim themselves -- this is a research brief,
    # not a black box. Headline sources are already linked inline above; this section is the
    # profile pages the stats tables came from.
    doc.add_heading("5. Sources", level=1)
    if player_url:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("FBref profile (season stats, bio, defensive/goalkeeping stats): ")
        _add_hyperlink(p, player_url, player_url)
    if xg and xg.get("understat_url"):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("Understat profile (xG/xA): ")
        _add_hyperlink(p, xg["understat_url"], xg["understat_url"])
    if articles:
        doc.add_paragraph(
            "News headlines are linked individually in \"What People Say\" above, each with "
            "its publisher and date.", style="List Bullet"
        )
    if not player_url and not (xg and xg.get("understat_url")) and not articles:
        doc.add_paragraph("No sourced links available for this brief.")

    # --- 6. Understanding the Signals --------------------------------------------------------
    # Standing, identical-every-run glossary -- not LLM-written, so it's the same trustworthy
    # explanation on every brief. Added per direct request (2026-09-17): a scout unfamiliar
    # with the tool shouldn't have to guess what "World Class" or "Hand-in-Glove Fit" means.
    doc.add_heading("6. Understanding the Signals", level=1)
    doc.add_paragraph().add_run(
        "Quality — a non-AI, percentile-based measure of this player's per-90 output against "
        "real players in the same league, season, and position group (450+ minutes played). "
        "Shown as both a raw percentile and a possession-adjusted one, since a player at a "
        "dominant, ball-hogging team naturally racks up fewer defensive actions than an equal "
        "player at a team that spends more time defending -- adjusting for that keeps a good "
        "defender at a big club from being penalized just for playing there."
    )
    quality_scale = doc.add_table(rows=0, cols=2)
    quality_scale.style = "Light Grid Accent 1"
    for cutoff, label in QUALITY_LABEL_SCALE:
        row = quality_scale.add_row().cells
        row[0].text = cutoff
        row[1].text = label
    doc.add_paragraph().add_run(
        "Quality does not know this player's transfer value, reputation, or what a scout has "
        "seen with their own eyes -- only what these specific numbers say relative to peers."
    ).italic = True

    ref_club_line = (
        f"For this brief, the reference is {fit_signal['reference_club']}'s current squad in "
        "the same position group."
        if fit_signal else
        "The reference club depends on which club philosophy is selected -- see NOTES.md or "
        "ask the tool's maintainer for the full six-club table."
    )
    doc.add_paragraph().add_run(
        "Fit (Hand-in-Glove Fit / Somewhat Fits / Completely Different) — measures how closely "
        "this player's statistical profile resembles the players who already play the chosen "
        f"club style, not how good the player is overall (that's Quality's job). {ref_club_line} "
        "A close match means this player produces value in the same specific ways that club's "
        "players do; a distant one doesn't mean the player is worse, just statistically "
        "differently shaped from that system's usual profile."
    )
    doc.add_paragraph().add_run(FIT_SCOPE_CAVEAT + " Treat Fit as one lens among several, not the deciding one.").italic = True

    doc.save(output_path)
    return output_path


def build_comparison_docx(
    candidates: list[dict],
    philosophy: dict | None,
    scout_notes: str | None,
    output_path: Path,
) -> Path:
    """Multi-candidate version of build_docx() -- one doc comparing 2+ players considered for
    the same role, instead of one doc per player. Each item in `candidates` is a research
    result dict with the same keys scoutlite_combined.research_player() returns (player_name,
    player_url, bio, stats, xg, articles, misc, keeper, news_synthesis, fit_read, quality,
    fit_signal, judge) -- built by a separate CLI (scoutlite_compare.py) that runs the exact
    same single-player pipeline once per candidate, so this function never recomputes a signal
    itself, only lays out results already computed elsewhere.

    Deliberately a separate function rather than build_docx() reused in a loop with different
    heading numbers: keeping the two independent avoids touching build_docx()'s existing,
    tested single-player output at all. Some content (news headline list) is trimmed per
    candidate versus the single-player brief, since a comparison with several full single-brief
    sections back to back gets unwieldy to skim -- the shared purpose here is picking a
    shortlist, not replacing each candidate's own full brief."""
    doc = Document()

    names = [c["player_name"] for c in candidates]
    title = doc.add_heading(f"Player Comparison — {', '.join(names)}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    has_philosophy = philosophy and (philosophy.get("in_possession") or philosophy.get("out_of_possession"))
    subtitle = doc.add_paragraph()
    subtitle_text = "A data/research layer for a scout to weigh -- not a scouting verdict."
    if has_philosophy:
        style_desc = " / ".join(v for v in philosophy.values() if v)
        subtitle_text = f"Compared against: {style_desc}  ·  " + subtitle_text
    subtitle.add_run(subtitle_text).italic = True
    if scout_notes and scout_notes.strip():
        doc.add_paragraph().add_run(f"Shared scout's role notes: \"{scout_notes.strip()}\"").italic = True

    # --- Summary table -----------------------------------------------------------------------
    doc.add_heading("Summary", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, label in enumerate(("Player", "Position", "Quality", "Fit", "Confidence")):
        header[i].text = label
    for c in candidates:
        row = table.add_row().cells
        row[0].text = c["player_name"]
        row[1].text = c["bio"].get("position", "unknown") if c.get("bio") else "unknown"
        quality = c.get("quality")
        row[2].text = quality["label"] if quality else "not available"
        fit_signal = c.get("fit_signal")
        row[3].text = f"{fit_signal['label']} vs. {fit_signal['reference_club']}" if fit_signal else "not available"
        judge = c.get("judge") or {}
        row[4].text = (
            f"{judge['source_accuracy']}%" + (" ⚠" if judge.get("confidence_warning") else "")
            if judge else "n/a"
        )
    doc.add_paragraph().add_run(
        "⚠ marks a candidate whose written sections scored below the automated confidence "
        "threshold -- check that candidate's own findings below before relying on the prose. "
        "The signals and data tables are unaffected either way."
    ).italic = True

    # --- Per-candidate sections ----------------------------------------------------------------
    for i, c in enumerate(candidates, start=1):
        doc.add_heading(f"Candidate {i}: {c['player_name']}", level=1)

        judge = c.get("judge") or {}
        if judge.get("confidence_warning"):
            warn = doc.add_paragraph()
            warn.add_run(
                f"⚠ CONFIDENCE WARNING — scored {judge['source_accuracy']}% on automated "
                f"claim-grounding after {judge['iterations']} revision pass(es), below the "
                f"{judge.get('threshold', 80)}% threshold."
            ).bold = True
            for finding in judge.get("findings", []):
                doc.add_paragraph(finding, style="List Bullet")

        quality = c.get("quality")
        fit_signal = c.get("fit_signal")
        quality_text = quality["label"] if quality else "not available"
        fit_text = f"{fit_signal['label']} vs. {fit_signal['reference_club']}" if fit_signal else "not available"
        doc.add_paragraph().add_run(f"Quality: {quality_text}  ·  Fit: {fit_text}").bold = True

        doc.add_heading("Who He Is", level=2)
        if c.get("bio"):
            _add_dict_table(doc, c["bio"])
        else:
            doc.add_paragraph("No background data available.")

        doc.add_heading("Stats & Performance", level=2)
        _add_dict_table(doc, c["stats"])
        if c.get("keeper"):
            _add_dict_table(doc, c["keeper"])
        if c.get("misc"):
            _add_dict_table(doc, c["misc"])
        if c.get("xg"):
            _add_dict_table(doc, c["xg"], skip=("understat_url",))

        doc.add_heading("What People Say", level=2)
        doc.add_paragraph(c.get("news_synthesis") or "No recent news synthesis available.")

        doc.add_heading("Signals & Fit Read", level=2)
        doc.add_paragraph(c.get("fit_read") or "No fit signal available.")

        doc.add_heading("Sources", level=2)
        if c.get("player_url"):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("FBref profile: ")
            _add_hyperlink(p, c["player_url"], c["player_url"])
        if c.get("xg") and c["xg"].get("understat_url"):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run("Understat profile: ")
            _add_hyperlink(p, c["xg"]["understat_url"], c["xg"]["understat_url"])

    # --- Understanding the Signals (shared, once) ----------------------------------------------
    doc.add_heading("Understanding the Signals", level=1)
    doc.add_paragraph().add_run(
        "Quality — a non-AI, percentile-based measure of per-90 output against real players in "
        "the same league, season, and position group (450+ minutes played)."
    )
    quality_scale = doc.add_table(rows=0, cols=2)
    quality_scale.style = "Light Grid Accent 1"
    for cutoff, label in QUALITY_LABEL_SCALE:
        row = quality_scale.add_row().cells
        row[0].text = cutoff
        row[1].text = label
    doc.add_paragraph().add_run(
        "Fit (Hand-in-Glove Fit / Somewhat Fits / Completely Different) — measures how closely "
        "a player's statistical profile resembles the players who already play the chosen club "
        "style, not how good the player is overall (that's Quality's job)."
    )
    doc.add_paragraph().add_run(FIT_SCOPE_CAVEAT).italic = True

    doc.save(output_path)
    return output_path
