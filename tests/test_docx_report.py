"""docx_report builds real .docx XML (hand-authored hyperlink runs) -- these guard against
the classic failure mode: XML that python-docx writes but Word can't open. Deterministic,
no network/LLM, so it belongs alongside the rest of the pure-layer suite."""
from docx import Document

from docx_report import FIT_SCOPE_CAVEAT, build_docx

STATS = {"season": "2023-2024", "squad": "Manchester City", "competition": "1. Premier League",
         "goals": "27", "minutes": "2,552"}
XG = {"understat_matched_name": "Erling Haaland", "understat_url": "https://understat.com/player/8260",
      "understat_team": "Manchester City", "xG": 28.8, "xA": 5.51}
ARTICLES = [
    {"title": "Haaland closing in on scoring record", "url": "https://example.com/a1",
     "source": {"name": "BBC Sport"}, "publishedAt": "2026-09-01T10:00:00Z"},
    {"title": "No-link wire report", "url": "", "source": {}, "publishedAt": ""},
]


def _build(tmp_path, **overrides):
    kwargs = dict(
        player_name="Erling Haaland", bio={"full_name": "Erling Braut Haaland"}, stats=STATS,
        xg=XG, articles=ARTICLES, misc=None, keeper=None,
        news_synthesis="No relevant news this window.", fit_read="Fit: 4/5 strong output.",
        scout_notes="tall, direct runner", philosophy={"in_possession": "vertical", "out_of_possession": ""},
        output_path=tmp_path / "brief.docx",
        quality={"score": 5, "position_group": "attack", "components": {"goals_per90": 91.0},
                 "avg_percentile": 91.0, "explanation": "Evaluated as an attacker..."},
        fit_score=4, judge={"source_accuracy": 95, "passed": True, "iterations": 1, "findings": [],
                             "confidence_warning": False, "threshold": 80},
        player_url="https://fbref.com/en/players/1f44ac21/Erling-Haaland",
    )
    kwargs.update(overrides)
    path = build_docx(**kwargs)
    return Document(path)


def _all_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_builds_and_reopens_without_error(tmp_path):
    doc = _build(tmp_path)
    assert "Erling Haaland — Player Research Brief" in _all_text(doc)


def test_headline_hyperlink_text_present(tmp_path):
    text = _all_text(_build(tmp_path))
    assert "Haaland closing in on scoring record" in text
    assert "BBC Sport" in text and "2026-09-01" in text


def test_headline_without_url_falls_back_to_plain_text(tmp_path):
    text = _all_text(_build(tmp_path))
    assert "No-link wire report" in text  # present, just not linked -- no crash on empty url


def test_sources_section_lists_fbref_and_understat_links(tmp_path):
    text = _all_text(_build(tmp_path))
    assert "5. Sources" in text
    assert "https://fbref.com/en/players/1f44ac21/Erling-Haaland" in text
    assert "https://understat.com/player/8260" in text


def test_sources_section_handles_nothing_available(tmp_path):
    text = _all_text(_build(tmp_path, xg=None, articles=[], player_url=None))
    assert "No sourced links available for this brief." in text


def test_understat_url_excluded_from_raw_stats_table(tmp_path):
    doc = _build(tmp_path)
    # the url gets its own hyperlinked "Source:" line, not a raw "Understat Url" table row
    table_cell_texts = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert "Understat Url" not in table_cell_texts


def test_fit_scope_caveat_shown_when_philosophy_given(tmp_path):
    # default fixture already sets a philosophy -- the caveat added after the Track C eval
    # (2026-09-16: 21/21 runs landed on Fit 3/5) must appear whenever Fit was actually assessed
    text = _all_text(_build(tmp_path))
    assert FIT_SCOPE_CAVEAT in text


def test_fit_scope_caveat_absent_when_no_philosophy_given(tmp_path):
    text = _all_text(_build(tmp_path, philosophy={"in_possession": "", "out_of_possession": ""}, fit_score=None))
    assert FIT_SCOPE_CAVEAT not in text


def test_specialist_caveat_rendered_when_present(tmp_path):
    quality = {"score": 4, "position_group": "attack",
               "components": {"goals_per90": 100.0, "assists_per90": 53.0},
               "avg_percentile": 76.5, "explanation": "Evaluated as an attacker...",
               "specialist_caveat": "This player's individual metrics vary widely..."}
    text = _all_text(_build(tmp_path, quality=quality))
    assert "vary widely" in text


def test_specialist_caveat_absent_when_not_flagged(tmp_path):
    # the default fixture's quality dict has no specialist_caveat key at all -- must not crash
    # on .get() and must not print anything about it
    text = _all_text(_build(tmp_path))
    assert "vary widely" not in text
